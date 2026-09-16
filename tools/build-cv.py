# Builds Mustafa Hadi's CV (ATS-safe, single column) from tools/cv-content.json.
# usage: python tools/build-cv.py OUT.docx [body_pt=10] [side_margin_in=0.45] [content.json]
#
# Layout order (recruiter-first): Header -> Summary -> Highlights -> Technical Skills ->
# Experience -> Leadership & Engineering Practices -> Projects -> Education -> Certifications.
# ATS safety: one column, no tables/text boxes/images, standard section names, plain ASCII
# punctuation (sanitized), real bullet characters, hyperlinks with visible text.
import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1]
BODY = float(sys.argv[2]) if len(sys.argv) > 2 else 10.0
MARGIN = float(sys.argv[3]) if len(sys.argv) > 3 else 0.45
CONTENT = sys.argv[4] if len(sys.argv) > 4 else os.path.join(HERE, "cv-content.json")

FONT = "Calibri"
TEXT_W = 8.5 - 2 * MARGIN
ACCENT = RGBColor(0x8B, 0x00, 0x00)
GREY = RGBColor(0x40, 0x40, 0x40)

with open(CONTENT, encoding="utf-8") as fh:
    C = json.load(fh)

# ── ATS sanitizer: plain ASCII punctuation only ─────────────────────────────
ASCII_MAP = {
    "→": " to ", "➔": " to ", "·": ", ", "•": "-", "–": "-", "—": "-",
    "‘": "'", "’": "'", "“": '"', "”": '"', "…": "...", " ": " ",
    "‑": "-", "−": "-", "×": "x", "≥": ">=", "≤": "<=", "≈": "~",
}


def ascii_clean(s):
    # Map typographic characters to ASCII; keep tabs and surrounding spaces intact
    # (tabs drive the right-aligned dates, and label runs end with ": ").
    for k, v in ASCII_MAP.items():
        s = s.replace(k, v)
    return s.encode("ascii", "ignore").decode("ascii")


# ── Document setup ──────────────────────────────────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(MARGIN)
sec.top_margin = sec.bottom_margin = Inches(max(MARGIN - 0.08, 0.3))

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(BODY)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.0
st.paragraph_format.widow_control = False

doc.core_properties.author = C["name"]
doc.core_properties.title = C["name"] + " - Resume"
doc.core_properties.subject = "Software Engineering Internship Resume"
doc.core_properties.keywords = ", ".join(C["skills"]["languages"] + C["skills"]["frameworks"])


def para(text="", size=None, bold=False, italic=False, align=None, before=0, after=0, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if text:
        run(p, text, bold=bold, italic=italic, size=size, color=color)
    return p


def run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(ascii_clean(text))
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return r


def hyperlink(p, url, text):
    r_id = p.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "1F4E79"); rPr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = ascii_clean(text); t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def rule(p, color="8B0000", size="6"):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", size), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def heading(text):
    p = para(ascii_clean(text).upper(), size=BODY + 1, bold=True, before=2, after=1, color=ACCENT)
    rule(p)
    return p


def right_tab(p):
    p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)


def entry(title, org, right, sub=None):
    p = para(before=3)
    right_tab(p)
    run(p, title, bold=True)
    if org:
        run(p, " - " + org, bold=True)
    run(p, "\t" + right, italic=True)
    if sub:
        para(sub, italic=True, color=GREY)


def bullet(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.17)
    pf.first_line_indent = Inches(-0.12)
    pf.space_after = Pt(0.2)
    r = p.add_run("•  ")  # real bullet glyph, parsed as a list marker by ATS
    run(p, text)
    return p


def labeled(label, text, after=0.5):
    p = para(after=after)
    run(p, label + ": ", bold=True)
    run(p, text)
    return p


# ── Header ──────────────────────────────────────────────────────────────────
para(C["name"].upper(), size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=ACCENT)
para(C["headline"], size=BODY + 0.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, f'{C["location"]}  |  {C["phone"]}  |  ')
hyperlink(p, "mailto:" + C["email"], C["email"])
run(p, "  |  " + C["work_auth"])
p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
for i, (label, url) in enumerate(C["links"]):
    if i:
        run(p, "  |  ")
    hyperlink(p, url, label)
if C.get("availability"):
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    run(p, "Availability: ", bold=True, color=GREY)
    run(p, C["availability"], italic=True, color=GREY)
else:
    p.paragraph_format.space_after = Pt(2)
rule(p, color="8B0000", size="8")

# ── Summary + highlights ────────────────────────────────────────────────────
heading("Professional Summary")
para(C["summary"], after=2)
if C.get("highlights"):
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    for i, h in enumerate(C["highlights"]):
        if i:
            run(p, "   |   ", color=ACCENT)
        run(p, h, bold=True)

# ── Technical skills (first thing after the summary) ────────────────────────
heading("Technical Skills")
S = C["skills"]


def joined(items):
    # Use semicolons when items themselves contain commas so ATS keyword splitting stays clean.
    return ("; " if any("," in x for x in items) else ", ").join(items)


labeled("Languages", joined(S["languages"]))
labeled("Frameworks & Libraries", joined(S["frameworks"]))
labeled("Cloud & Infrastructure", joined(S["cloud"]))
labeled("Data, APIs & Tools", joined(S["data_tools"]))
labeled("Engineering Practices", joined(S["practices"]))

# ── Experience ──────────────────────────────────────────────────────────────
heading("Experience")
for e in C["experience"]:
    entry(e["title"], e["company"], f'{e["location"]}  |  {e["period"]}', e.get("subline"))
    for b in e["bullets"]:
        bullet(b)

# ── Leadership & practices ──────────────────────────────────────────────────
heading(C["leadership"]["heading"])
for b in C["leadership"]["bullets"]:
    bullet(b)

# ── Projects ────────────────────────────────────────────────────────────────
heading("Projects")
for pr in C["projects"]:
    p = para(before=3)
    right_tab(p)
    run(p, pr["name"], bold=True)
    run(p, " - " + pr["stack"] + "\t")
    hyperlink(p, "https://" + pr["url"], pr["url"])
    for b in pr["bullets"]:
        bullet(b)

# ── Education + certifications ──────────────────────────────────────────────
heading("Education & Certifications")
ed = C["education"]
entry(ed["school"], ed["degree"], ed["right"])
if ed.get("note"):
    para(ed["note"])
for c in C["certifications"]:
    bullet(c)

# ── Modern Word compatibility (no "Compatibility Mode" banner) ──────────────
settings = doc.settings.element
compat = settings.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat")
    settings.append(compat)
for cs in list(compat.findall(qn("w:compatSetting"))):
    if cs.get(qn("w:name")) == "compatibilityMode":
        compat.remove(cs)
cs = OxmlElement("w:compatSetting")
cs.set(qn("w:name"), "compatibilityMode")
cs.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
cs.set(qn("w:val"), "15")
compat.append(cs)

doc.save(OUT)
print("saved", OUT)
